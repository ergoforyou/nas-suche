from __future__ import annotations

import io
import struct
import time
from email.message import EmailMessage
from pathlib import Path

import pymupdf
import pytest
from PIL import Image

from agentur_pdf.kern.auftraege import Auftrag
from agentur_pdf.kern.dokumente import DokumentFehler
from agentur_pdf.kern.konvertierung import html_bereinigen, konvertieren, libreoffice_finden
from agentur_pdf.kern.outlook_msg import filetime_zu_datum, zeit_aus_eigenschaften
from agentur_pdf.kern.zusammenfuehren import Eintrag, umwandeln, zusammenfuehren

mit_libreoffice = pytest.mark.skipif(libreoffice_finden() is None, reason="LibreOffice nicht installiert")


def text(pdf: bytes) -> list[str]:
    return [s.get_text().strip() for s in pymupdf.open("pdf", pdf)]


def bild(pfad: Path, groesse=(400, 300), farbe=(200, 30, 30), **speichern) -> Path:
    Image.new("RGB", groesse, farbe).save(pfad, **speichern)
    return pfad


def docx(pfad: Path, absatz: str) -> Path:
    import docx as python_docx

    d = python_docx.Document()
    d.add_heading("Versicherungsangebot", 1)
    d.add_paragraph(absatz)
    d.save(pfad)
    return pfad


def test_pdf_wird_durchgereicht(brief):
    assert len(text(konvertieren("brief.pdf", pfad=str(brief)).pdf)) == 3


def test_verschluesselte_pdf_abgelehnt(tmp_path):
    p = tmp_path / "g.pdf"
    d = pymupdf.open()
    d.new_page()
    d.save(p, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="x")
    with pytest.raises(DokumentFehler, match="passwortgeschützt"):
        konvertieren("g.pdf", pfad=str(p))


def test_jpg_mit_exif_drehung(tmp_path):
    p = tmp_path / "foto.jpg"
    img = Image.new("RGB", (400, 200), (10, 10, 10))
    exif = img.getexif()
    exif[0x0112] = 6  # 90° im Uhrzeigersinn → Hochformat
    img.save(p, exif=exif.tobytes())
    seite = pymupdf.open("pdf", konvertieren("foto.jpg", pfad=str(p)).pdf)[0]
    assert seite.rect.height > seite.rect.width


def test_riesiges_foto_wird_auf_a4_gebracht(tmp_path):
    p = bild(tmp_path / "handy.png", (4000, 3000))
    seite = pymupdf.open("pdf", konvertieren("handy.png", pfad=str(p)).pdf)[0]
    assert seite.rect.width <= 842.5 and seite.rect.height <= 842.5
    assert abs(seite.rect.width / seite.rect.height - 4 / 3) < 0.01  # Seitenverhältnis bleibt


def test_mehrseitiges_tiff(tmp_path):
    p = tmp_path / "scan.tif"
    seiten = [Image.new("L", (850, 1100), v) for v in (255, 200, 100)]
    seiten[0].save(p, save_all=True, append_images=seiten[1:], dpi=(100, 100))
    doc = pymupdf.open("pdf", konvertieren("scan.tif", pfad=str(p)).pdf)
    assert doc.page_count == 3
    assert abs(doc[0].rect.width - 850 * 72 / 100) < 1  # Auflösung wird beachtet


def test_heic(tmp_path):
    pillow_heif = pytest.importorskip("pillow_heif")
    p = tmp_path / "iphone.heic"
    pillow_heif.from_pillow(Image.new("RGB", (300, 400), (0, 120, 0))).save(p)
    doc = pymupdf.open("pdf", konvertieren("iphone.heic", pfad=str(p)).pdf)
    assert doc.page_count == 1 and doc[0].get_images()


def test_text_mit_umlauten_cp1252(tmp_path):
    p = tmp_path / "notiz.txt"
    p.write_bytes("Prämie für Müller: 312,40 € – Grüße".encode("cp1252"))
    assert "Prämie für Müller: 312,40 € – Grüße" in text(konvertieren("notiz.txt", pfad=str(p)).pdf)[0]


def test_eml_mit_anhaengen(tmp_path, brief):
    m = EmailMessage()
    m["Subject"] = "Schadenmeldung Müller"
    m["From"] = "Kunde <kunde@example.de>"
    m["To"] = "agentur@example.de"
    m["Date"] = "Tue, 13 Oct 2026 09:15:00 +0200"
    m.set_content("Sehr geehrte Damen und Herren,\nanbei die Fotos.\nGrüße")
    m.add_attachment(brief.read_bytes(), maintype="application", subtype="pdf", filename="Police.pdf")
    puffer = io.BytesIO()
    Image.new("RGB", (100, 100), (1, 2, 3)).save(puffer, "PNG")
    m.add_attachment(puffer.getvalue(), maintype="image", subtype="png", filename="Foto.png")
    m.add_attachment(b"PK..", maintype="application", subtype="zip", filename="daten.zip")
    p = tmp_path / "mail.eml"
    p.write_bytes(bytes(m))

    teil = konvertieren("mail.eml", pfad=str(p))
    kopf = text(teil.pdf)[0]
    assert "Schadenmeldung Müller" in kopf and "kunde@example.de" in kopf
    assert "13.10.2026" in kopf and "daten.zip" in kopf  # nicht umwandelbarer Anhang wird genannt
    assert [u.titel for u in teil.unterteile] == ["Police", "Foto"]
    assert konvertieren("mail.eml", pfad=str(p), anhaenge=False).unterteile == []


def test_html_mail_wird_bereinigt():
    roh = "<html><head><style>x{}</style></head><body><script>alert(1)</script><p>Hallo</p><img src='http://x/t.gif'></body></html>"
    assert html_bereinigen(roh) == "<p>Hallo</p>"


def test_msg_hilfsfunktionen():
    # 13.10.2026 07:15 UTC als FILETIME
    wert = int((time.mktime((2026, 10, 13, 7, 15, 0, 0, 0, 0)) - time.timezone + 11644473600) * 10_000_000)
    datum = filetime_zu_datum(wert)
    assert (datum.year, datum.month, datum.day) == (2026, 10, 13)
    props = bytes(32) + struct.pack("<HHIq", 0x0040, 0x0039, 0, wert)
    assert zeit_aus_eigenschaften(props, 0x0039, 32) == datum


def test_unbekannter_typ():
    with pytest.raises(DokumentFehler, match="nicht unterstützt"):
        konvertieren("archiv.zip", daten=b"PK")


@mit_libreoffice
def test_docx_und_xlsx(tmp_path):
    import openpyxl

    t = konvertieren("angebot.docx", pfad=str(docx(tmp_path / "angebot.docx", "Jahresprämie 312,40 € für Hausrat")))
    assert "Jahresprämie 312,40 € für Hausrat" in text(t.pdf)[0]
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Beitragsübersicht"
    wb.active["B2"] = 1234.5
    wb.save(tmp_path / "tabelle.xlsx")
    assert "Beitragsübersicht" in text(konvertieren("tabelle.xlsx", pfad=str(tmp_path / "tabelle.xlsx")).pdf)[0]


@mit_libreoffice
def test_t1_pdf_jpg_docx_zusammenfuehren(speicher, tmp_path, brief):
    """T1: PDF, JPG, DOCX → eine PDF in richtiger Reihenfolge, in max. 10 s."""
    jpg = bild(tmp_path / "foto.jpg", (600, 800), quality=90)
    wort = docx(tmp_path / "anschreiben.docx", "Anschreiben an Familie Weiß")
    eintraege = [Eintrag(p.name, str(p)) for p in (brief, jpg, wort)]
    start = time.perf_counter()
    ergebnis = zusammenfuehren(eintraege, speicher, Auftrag("t", "T1"), lambda f, *a, **k: f(*a, **k))
    dauer = time.perf_counter() - start
    assert dauer < 10, f"{dauer:.1f} s"
    dok = speicher.get(ergebnis["dokumente"][0])
    seiten = [s.get_text() for s in dok.doc]
    assert len(seiten) == 5  # 3 PDF + 1 Bild + 1 Word
    assert "Versicherungsschein" in seiten[0]
    assert dok.doc[3].get_images() and not seiten[3].strip()  # Bild an Position 4
    assert "Familie Weiß" in seiten[4]
    assert [t[1:] for t in dok.doc.get_toc()] == [["brief", 1], ["Abschnitt 1", 1], ["Abschnitt 2", 2], ["Abschnitt 3", 3], ["foto", 4], ["anschreiben", 5], ["Versicherungsangebot", 5]]
    assert dok.doc.get_toc()[-1][0] == 2  # Word-Überschrift als Unter-Lesezeichen
    assert dok.titel == "brief (zusammengeführt).pdf" and dok.geaendert
    assert dok.ordner_vorschlag == str(tmp_path)


def test_zusammenfuehren_ohne_lesezeichen(speicher, tmp_path, brief):
    png = bild(tmp_path / "a.png")
    e = [Eintrag("brief.pdf", str(brief)), Eintrag("a.png", str(png))]
    ergebnis = zusammenfuehren(e, speicher, Auftrag("t", "x"), lambda f, *a, **k: f(*a, **k), lesezeichen=False)
    assert speicher.get(ergebnis["dokumente"][0]).doc.get_toc() == []


def test_umwandeln_speichert_neben_quelle(speicher, tmp_path):
    p = tmp_path / "notiz.txt"
    p.write_text("Hallo", encoding="utf-8")
    (tmp_path / "notiz.pdf").write_bytes(b"alt")  # Konflikt → „notiz (2).pdf“
    ergebnis = umwandeln([Eintrag("notiz.txt", str(p))], speicher, Auftrag("t", "x"), lambda f, *a, **k: f(*a, **k), speichern=True)
    assert ergebnis["dateien"] == [str(tmp_path / "notiz (2).pdf")]
    dok = speicher.get(ergebnis["dokumente"][0])
    assert dok.pfad == str(tmp_path / "notiz (2).pdf") and not dok.geaendert


def test_abbrechen(speicher, tmp_path, brief):
    from agentur_pdf.kern.auftraege import Abgebrochen

    a = Auftrag("t", "x")
    a.abbrechen()
    with pytest.raises(Abgebrochen):
        zusammenfuehren([Eintrag("brief.pdf", str(brief))], speicher, a, lambda f, *x, **k: f(*x, **k))


def test_adressen_lesbar():
    from agentur_pdf.kern.konvertierung import adressen_lesbar

    assert adressen_lesbar('"\'werkstatt@example.de\'" <werkstatt@example.de>') == "werkstatt@example.de"
    assert adressen_lesbar('"Müller, Anna (AV)" <anna@example.de>, b@example.de') == "Müller, Anna (AV) <anna@example.de>, b@example.de"
