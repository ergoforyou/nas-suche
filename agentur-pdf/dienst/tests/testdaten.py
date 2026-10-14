"""Erzeugt synthetische Testdokumente (keine echten Kundendaten)."""

from __future__ import annotations

from pathlib import Path

import pymupdf

BRIEF = """Sehr geehrte Frau Müller,

anbei erhalten Sie Ihren neuen Versicherungsschein zur Hausratversicherung.
Die jährliche Prämie beträgt ab dem 01.01.2027 insgesamt 312,40 EUR.
Bitte prüfen Sie die Angaben zur Wohnfläche und zum Gebäude in der Straße.
Bei Fragen erreichen Sie Ihre Agentur werktags von 8 bis 17 Uhr.

Mit freundlichen Grüßen
Ihre Versicherungsagentur"""


def versicherungsbrief(pfad: Path, seiten: int = 3, mit_lesezeichen: bool = True) -> Path:
    doc = pymupdf.open()
    for nr in range(seiten):
        seite = doc.new_page(width=595, height=842)
        seite.insert_text((72, 80), f"Versicherungsschein Nr. VS-{1000 + nr}", fontsize=16, fontname="helv")
        seite.insert_textbox(pymupdf.Rect(72, 110, 523, 500), BRIEF, fontsize=11, fontname="helv")
        seite.insert_text((280, 810), f"Seite {nr + 1} von {seiten}", fontsize=9)
        seite.draw_rect(pymupdf.Rect(72, 520, 523, 700), color=(0.7, 0.1, 0.1), width=1)
    if mit_lesezeichen:
        doc.set_toc([[1, f"Abschnitt {i + 1}", i + 1] for i in range(seiten)])
    doc.save(pfad)
    doc.close()
    return pfad


def gedrehtes_dokument(pfad: Path) -> Path:
    doc = pymupdf.open()
    seite = doc.new_page(width=595, height=842)
    seite.insert_text((72, 100), "Prämie quer", fontsize=20)
    seite.set_rotation(90)
    doc.save(pfad)
    doc.close()
    return pfad


def beschriftet(pfad: Path, praefix: str, anzahl: int) -> Path:
    """Seiten mit großer Beschriftung A1, A2 … (gut prüfbar per Suche)."""
    doc = pymupdf.open()
    for i in range(anzahl):
        doc.new_page(width=595, height=842).insert_text((72, 120), f"{praefix}{i + 1}", fontsize=48)
    doc.save(pfad)
    doc.close()
    return pfad


def mit_leerseiten(pfad: Path) -> Path:
    """Duplex-Scan-Nachbildung: Inhaltsseiten mit leeren Rückseiten."""
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page().insert_text((72, 120), f"Inhalt {i + 1}", fontsize=24)
        doc.new_page()  # leere Rückseite
    doc.save(pfad)
    doc.close()
    return pfad


def foto(pfad: Path) -> Path:
    from PIL import Image

    Image.new("RGB", (600, 800), (180, 40, 40)).save(pfad, quality=90)
    return pfad


def word(pfad: Path, absatz: str) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Versicherungsangebot", 1)
    d.add_paragraph(absatz)
    d.save(pfad)
    return pfad
